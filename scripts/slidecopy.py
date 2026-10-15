# -*- coding: utf-8 -*-
"""跨文件整页复制（含图片等关系部件）、文本置换、安全删页。作为模块导入使用：

    from pptx import Presentation
    from slidecopy import copy_slide, replace_map, delete_slide
    dst, src = Presentation("mine.pptx"), Presentation("reference.pptx")
    s = copy_slide(dst, src.slides[42], index=47)
    replace_map(s, {"对照组 50 名": "医护端 ≥3 名"})
    dst.save("mine - 副本.pptx")

复制后务必全篇搜参考稿的关键词（零命中），并核对位图里的文字与图标叙事。"""
import copy

from pptx.oxml.ns import qn

REL_ATTRS = [qn("r:embed"), qn("r:link"), qn("r:id"), qn("r:pict"), qn("r:dm"), qn("r:lo"), qn("r:qs"), qn("r:cs")]


def _pick_layout(dst_prs, src_slide):
    want = src_slide.slide_layout.name
    for lay in dst_prs.slide_layouts:
        if lay.name == want:
            return lay
    for lay in dst_prs.slide_layouts:
        if lay.name in ("仅标题", "Title Only"):
            return lay
    return dst_prs.slide_layouts[-1]


def _remap_rels(spTree, src_part, dst_part):
    cache = {}
    for el in spTree.iter():
        for a in REL_ATTRS:
            rid = el.get(a)
            if not rid:
                continue
            if rid not in cache:
                rel = src_part.rels.get(rid)
                if rel is None:
                    el.attrib.pop(a, None)
                    continue
                if rel.is_external:
                    cache[rid] = dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
                else:
                    cache[rid] = dst_part.relate_to(rel.target_part, rel.reltype)
            el.set(a, cache[rid])


def renumber_ids(slide):
    """复制来的形状沿用源 cNvPr/@id，会造成同页重复 id，统一重编。"""
    for i, el in enumerate(slide.shapes._spTree.iter(qn("p:cNvPr")), start=2):
        el.set("id", str(i))


def copy_slide(dst_prs, src_slide, index=None):
    """把 src_slide 复制进 dst_prs，返回新 slide。index 为 0 基插入位置。"""
    new = dst_prs.slides.add_slide(_pick_layout(dst_prs, src_slide))
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    spTree = new.shapes._spTree
    for el in src_slide.shapes._spTree:
        if el.tag.endswith("}nvGrpSpPr") or el.tag.endswith("}grpSpPr"):
            continue
        spTree.append(copy.deepcopy(el))
    _remap_rels(spTree, src_slide.part, new.part)
    renumber_ids(new)
    if index is not None:
        lst = dst_prs.slides._sldIdLst
        last = list(lst)[-1]
        lst.remove(last)
        lst.insert(index, last)
    return new


def delete_slide(prs, index):
    """先解除关系再摘条目；只摘 sldIdLst 会留下孤儿部件，保存时报 Duplicate name。"""
    lst = prs.slides._sldIdLst
    sld = list(lst)[index]
    prs.part.drop_rel(sld.rId)
    lst.remove(sld)


def set_text(shape, lines):
    """把形状文本换成 lines（列表 = 多段），保留原第一段第一个 run 的格式。"""
    tf = shape.text_frame
    if not tf.paragraphs or not tf.paragraphs[0].runs:
        tf.text = "\n".join(lines)
        return
    proto = copy.deepcopy(tf.paragraphs[0]._p)
    parent = tf.paragraphs[0]._p.getparent()
    for p in parent.findall(qn("a:p")):
        parent.remove(p)
    for ln in lines:
        np_ = copy.deepcopy(proto)
        runs = np_.findall(qn("a:r"))
        for r in runs[1:]:
            np_.remove(r)
        runs[0].find(qn("a:t")).text = ln
        parent.append(np_)


def _text_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from _text_shapes(sh.shapes)
        elif sh.has_text_frame and sh.text_frame.text.strip():
            yield sh


def replace_map(slide, mapping):
    """按 {旧整段文本: 新文本或行列表} 置换；打印未命中的键和未被置换的残留文本。"""
    hit = set()
    for sh in _text_shapes(slide.shapes):
        cur = sh.text_frame.text.strip().replace("\n", "")
        key = next((k for k in mapping if k.replace("\n", "") == cur), None)
        if key is None:
            print("   [未替换] " + cur[:50])
            continue
        v = mapping[key]
        set_text(sh, list(v) if isinstance(v, (list, tuple)) else [v])
        hit.add(key)
    for k in mapping:
        if k not in hit:
            print("   [未命中] " + k[:40])
    return hit

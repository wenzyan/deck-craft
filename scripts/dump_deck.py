# -*- coding: utf-8 -*-
"""逐页导出文本与几何（pt），每轮动手前跑一次。用法: python dump_deck.py <pptx> [页码...]"""
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.stdout.reconfigure(encoding="utf-8")
PT = 12700.0


def walk(shapes, depth=0):
    for sh in shapes:
        geo = ""
        if sh.width is not None:
            geo = "L%.0f T%.0f W%.0f H%.0f" % (sh.left / PT, sh.top / PT, sh.width / PT, sh.height / PT)
        text = sh.text_frame.text.replace("\n", " | ")[:80] if sh.has_text_frame else ""
        kind = "GROUP" if sh.shape_type == MSO_SHAPE_TYPE.GROUP else ("PIC" if sh.shape_type == MSO_SHAPE_TYPE.PICTURE else "")
        print("%s- %s %s %s %s" % ("  " * depth, sh.name, kind, geo, text))
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            walk(sh.shapes, depth + 1)  # 组内坐标是相对组的，不是页面绝对坐标


def main():
    prs = Presentation(sys.argv[1])
    pages = {int(x) for x in sys.argv[2:]}
    print("画布 %.0f×%.0f pt，共 %d 页" % (prs.slide_width / PT, prs.slide_height / PT, len(prs.slides)))
    for n, sl in enumerate(prs.slides, start=1):
        if pages and n not in pages:
            continue
        print("\n=== P%d [%s]" % (n, sl.slide_layout.name))
        walk(sl.shapes)
        if sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip():
            print("  备注: " + sl.notes_slide.notes_text_frame.text[:120].replace("\n", " "))


if __name__ == "__main__":
    main()

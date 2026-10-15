# -*- coding: utf-8 -*-
"""从 JSON 页面描述生成原生可编辑 pptx。用法: python deck.py <spec.json> <out.pptx>
版式为纵向流式排布：块按 y 游标依次放置，结构上不会重叠；放不下时打印警告（该拆页了）。
spec 格式见 references/spec-format.md。"""
import json
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

sys.stdout.reconfigure(encoding="utf-8")

THEME = {
    "accent": "9D2933", "text": "1A1A1A", "grey": "595959", "line": "BFBFBF", "title": "1A1A1A",
    "font_ea": "微软雅黑", "font_latin": "Arial", "size_body": 14, "size_title": 28,
}
SW, SH = 960, 540            # 画布 pt（16:9）
ML, MR = 48, 48              # 左右边距
BODY_T, BODY_B = 92, 500     # 正文区上下边界
GAP = 12                     # 块间距


def rgb(key):
    return RGBColor.from_string(THEME.get(key, key))


def wid(s):
    """估算字宽（以字号为单位）：全角 1，半角 0.55。"""
    return sum(1.0 if ord(c) > 0x2E80 else 0.55 for c in s)


def fit_size(lines, width, hi=None, lo=11):
    """同一组文字取同一字号：最长一行能单行放下的最大字号。"""
    hi = hi or THEME["size_body"]
    longest = max((wid(t) for t in lines), default=1) or 1
    return max(lo, min(hi, int((width - 12) / longest)))


# ---------- 底层构件 ----------

def add_run(p, text, size, bold=False, color="text"):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size, f.bold, f.color.rgb = Pt(size), bold, rgb(color)
    f.name = THEME["font_latin"]
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        etree.SubElement(rPr, qn(tag)).set("typeface", THEME["font_ea"])
    return r


def box(slide, x, y, w, h, fill=None, line=None, lw=0.75, prst=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(prst, Pt(x), Pt(y), Pt(w), Pt(h))
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb, sp.line.width = rgb(line), Pt(lw)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    tf = sp.text_frame
    tf.word_wrap, tf.vertical_anchor = True, MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Pt(6)
    tf.margin_top = tf.margin_bottom = Pt(3)
    return sp


def text(slide, x, y, w, h, lines, size, bold=False, color="text", align=PP_ALIGN.LEFT,
         anchor=None, shape=None):
    """写文字；shape 给定时写进该形状（默认垂直居中），否则新建文本框（默认顶对齐）。"""
    tf = (shape or slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))).text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor or (MSO_ANCHOR.MIDDLE if shape is not None else MSO_ANCHOR.TOP)
    if shape is None:
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = [lines] if isinstance(lines, str) else lines
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(size * 0.35 if i < len(lines) - 1 else 0)
        add_run(p, ln, size, bold, color)
    return tf


def hline(slide, x, y, w, color="accent", lw=1.5):
    c = slide.shapes.add_connector(1, Pt(x), Pt(y), Pt(x + w), Pt(y))
    c.line.color.rgb, c.line.width = rgb(color), Pt(lw)


# ---------- 块：每个函数 (slide, x, y, w, spec) -> 占用高度 ----------

def b_banner(s, x, y, w, v):
    size = fit_size([v], w - 24, hi=17)
    sp = box(s, x, y, w, size * 2.2, line="accent", lw=1.25)
    text(s, 0, 0, 0, 0, v, size, True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=sp)
    return size * 2.2


def b_cards(s, x, y, w, v):
    n, gap = len(v), 16
    cw = (w - gap * (n - 1)) / n
    size = fit_size([ln for c in v for ln in c["lines"]], cw - 12)
    rows = max(len(c["lines"]) for c in v)
    hh, bh = size * 2.2, rows * size * 1.45 + 16
    for i, c in enumerate(v):
        cx = x + i * (cw + gap)
        text(s, cx, y, cw, hh - 4, c["head"], size + 2, True, anchor=MSO_ANCHOR.BOTTOM)
        hline(s, cx, y + hh, cw)
        sp = box(s, cx, y + hh + 4, cw, bh, line="line")
        text(s, 0, 0, 0, 0, c["lines"], size, shape=sp, anchor=MSO_ANCHOR.TOP)
    return hh + 4 + bh


def b_grid(s, x, y, w, v):
    rows, cols, lw, gap = v["rows"], v["cols"], 72, 10
    cw = (w - lw - gap * len(cols)) / len(cols)
    size = fit_size([c for col in cols for c in col["cells"]], cw - 12)
    hh, rh = size * 2.2, size * 2.6
    for j, col in enumerate(cols):
        cx = x + lw + gap + j * (cw + gap)
        text(s, cx, y, cw, hh - 4, col["head"], size + 1, True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
        hline(s, cx, y + hh, cw)
    for i, name in enumerate(rows):
        ry = y + hh + 6 + i * (rh + 6)
        lab = box(s, x, ry, lw, rh, fill="accent")
        text(s, 0, 0, 0, 0, name, size, True, color="FFFFFF", align=PP_ALIGN.CENTER, shape=lab)
        for j, col in enumerate(cols):
            sp = box(s, x + lw + gap + j * (cw + gap), ry, cw, rh, line="line")
            text(s, 0, 0, 0, 0, col["cells"][i], size, align=PP_ALIGN.CENTER, shape=sp)
    return hh + 6 + len(rows) * (rh + 6) - 6


def b_bignums(s, x, y, w, v):
    n = len(v)
    cw = w / n
    for i, it in enumerate(v):
        cx = x + i * cw
        text(s, cx, y, cw, 46, it["num"], 36, True, color="accent", align=PP_ALIGN.CENTER)
        text(s, cx, y + 50, cw, 22, it["label"], 13, color="grey", align=PP_ALIGN.CENTER)
    return 74


def b_chain(s, x, y, w, v):
    n, aw = len(v), 26
    bw = (w - aw * (n - 1) - 8 * (n - 1)) / n
    size = fit_size(v, bw - 8, hi=15)
    bh = size * 2.6
    for i, t in enumerate(v):
        bx = x + i * (bw + aw + 8)
        sp = box(s, bx, y, bw, bh, line="accent", lw=1.25, prst=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(s, 0, 0, 0, 0, t, size, True, align=PP_ALIGN.CENTER, shape=sp)
        if i < n - 1:
            ah = min(26, max(9, bh * 0.4))   # 箭头粗细按方框高的 40%，限制在 9–26pt
            box(s, bx + bw + 4, y + (bh - ah) / 2, aw, ah, fill="accent", prst=MSO_SHAPE.RIGHT_ARROW)
    return bh


def b_figure(s, x, y, w, v):
    h, path, cap = v.get("h", 200), v.get("path"), v.get("caption", "")
    if path:
        pic = s.shapes.add_picture(path, Pt(x), Pt(y), height=Pt(h))
        if pic.width > Pt(w):
            pic.width, pic.height = Pt(w), int(pic.height * Pt(w) / pic.width)
        pic.left = int(Pt(x) + (Pt(w) - pic.width) / 2)
        h = pic.height / 12700
    else:
        sp = box(s, x, y, w, h, fill="F2F2F2", line="line")
        sp.line.dash_style = 4
        text(s, 0, 0, 0, 0, "图位：" + v.get("todo", cap), 12, color="grey", align=PP_ALIGN.CENTER, shape=sp)
    if cap:
        text(s, x, y + h + 4, w, 18, cap, 11, color="grey", align=PP_ALIGN.CENTER)
        h += 22
    return h


def b_bullets(s, x, y, w, v):
    size = THEME["size_body"]
    rows = sum(max(1, int(wid(t) * size / (w - 16)) + 1) for t in v)
    text(s, x, y, w, rows * size * 1.5, ["▪ " + t for t in v], size)
    return rows * size * 1.5


def b_note(s, x, y, w, v):
    text(s, x, y, w, 16, v, 10, color="grey")
    return 16


def b_split(s, x, y, w, v):
    ratio, gap = v.get("ratio", 0.5), 24
    lw = (w - gap) * ratio
    hl = flow(s, x, y, lw, v["left"])
    hr = flow(s, x + lw + gap, y, w - gap - lw, v["right"])
    return max(hl, hr)


BLOCKS = {"banner": b_banner, "cards": b_cards, "grid": b_grid, "bignums": b_bignums, "chain": b_chain,
          "figure": b_figure, "bullets": b_bullets, "note": b_note, "split": b_split}


def flow(s, x, y, w, blocks):
    y0 = y
    for blk in blocks:
        (kind, v), = blk.items()
        y += BLOCKS[kind](s, x, y, w, v) + GAP
    return y - y0 - GAP


# ---------- 页面 ----------

def page_number(s):
    tf = s.shapes.add_textbox(Pt(SW - 80), Pt(SH - 30), Pt(50), Pt(18)).text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    fld = etree.SubElement(p._p, qn("a:fld"), id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
    rpr = etree.SubElement(fld, qn("a:rPr"), lang="zh-CN", sz="1100")
    etree.SubElement(etree.SubElement(rpr, qn("a:solidFill")), qn("a:srgbClr"), val=THEME["grey"])
    etree.SubElement(fld, qn("a:t")).text = "‹#›"


def content(s, pg, n):
    text(s, ML, 30, SW - ML - MR, 44, pg["title"], THEME["size_title"], True, color="title", anchor=MSO_ANCHOR.BOTTOM)
    hline(s, ML, 80, SW - ML - MR, color="line", lw=0.75)
    used = flow(s, ML, BODY_T, SW - ML - MR, pg.get("blocks", []))
    if BODY_T + used > BODY_B:
        print("警告 P%d「%s」超出正文区 %.0fpt：删减或拆页" % (n, pg["title"], BODY_T + used - BODY_B))
    elif used < 0.7 * (BODY_B - BODY_T):
        print("提示 P%d「%s」只用了正文区 %.0f%%：从依据补内容、放大视觉主体或与相邻页合并"
              % (n, pg["title"], 100 * used / (BODY_B - BODY_T)))


def cover(s, pg, n):
    box(s, 0, 0, 14, SH, fill="accent")
    text(s, 80, 170, SW - 160, 100, pg["title"], 36, True, color="title", anchor=MSO_ANCHOR.BOTTOM)
    hline(s, 80, 282, 120, lw=3)
    text(s, 80, 296, SW - 160, 120, [pg.get("subtitle", "")] + pg.get("meta", []), 16, color="grey")


def section(s, pg, n):
    box(s, 0, 0, SW, SH, fill="accent")
    text(s, 80, 180, 200, 90, pg.get("no", ""), 60, True, color="FFFFFF")
    text(s, 80, 280, SW - 160, 60, pg["title"], 32, True, color="FFFFFF")


def end(s, pg, n):
    text(s, 0, 220, SW, 80, pg.get("title", "谢谢"), 40, True, color="title", align=PP_ALIGN.CENTER)


PAGES = {"content": content, "cover": cover, "section": section, "end": end}


def build(spec, out):
    THEME.update(spec.get("theme", {}))
    prs = Presentation(spec["template"]) if spec.get("template") else Presentation()
    if (prs.slide_width, prs.slide_height) != (Pt(SW), Pt(SH)):
        if spec.get("template"):
            sys.exit("模板须为 16:9（960×540pt），当前 %.0f×%.0fpt" % (prs.slide_width / 12700, prs.slide_height / 12700))
        prs.slide_width, prs.slide_height = Pt(SW), Pt(SH)
    layout = next((l for l in prs.slide_layouts if l.name in ("Blank", "空白")), prs.slide_layouts[-1])
    for n, pg in enumerate(spec["slides"], start=1):
        s = prs.slides.add_slide(layout)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        PAGES[pg.get("type", "content")](s, pg, n)
        if pg.get("type", "content") == "content":
            page_number(s)
        if pg.get("notes"):
            s.notes_slide.notes_text_frame.text = pg["notes"]
    prs.save(out)
    print("已生成 %s，%d 页" % (out, len(spec["slides"])))


if __name__ == "__main__":
    build(json.load(open(sys.argv[1], encoding="utf-8")), sys.argv[2])

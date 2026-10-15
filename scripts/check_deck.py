# -*- coding: utf-8 -*-
"""几何核验：越界、重叠、文字溢出。用法: python check_deck.py <pptx> [页码...]"""
import math
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
PT = 12700.0


def wid(s):
    """估算字宽（以字号为单位）：全角 1，半角 0.55。"""
    return sum(1.0 if ord(c) > 0x2E80 else 0.55 for c in s)


def skip_fit(sh):
    tf = sh.text_frame
    body = tf._txBody.find(qn("a:bodyPr"))
    vert = body is not None and body.get("vert") not in (None, "horz")
    return vert or tf.word_wrap is False or tf.auto_size == MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT


def need_height(sh, w):
    need, avail = 0.0, max(w - 15, 1)
    for p in sh.text_frame.paragraphs:
        if not p.runs:
            continue
        sizes = [r.font.size.pt if r.font.size else 18 for r in p.runs]
        line_w = sum(wid(r.text) * s for r, s in zip(p.runs, sizes))
        rows = max(1, math.ceil(line_w / avail))
        need += rows * max(sizes) * 1.16 + (p.space_after.pt if p.space_after else 0)
    return need


def contains(a, b):
    """a 完整包住 b（文字压在底框上属有意叠放）。"""
    return a[1] - 1 <= b[1] and a[2] - 1 <= b[2] and b[1] + b[3] <= a[1] + a[3] + 1 and b[2] + b[4] <= a[2] + a[4] + 1


def check_slide(n, sl, W, H):
    issues, rects = [], []
    for sh in sl.shapes:
        if sh.width is None:
            continue
        x, y, w, h = sh.left / PT, sh.top / PT, sh.width / PT, sh.height / PT
        rects.append((sh.name, x, y, w, h, w > W * 0.97 and h > H * 0.97))
        if x < -1 or y < -1 or x + w > W + 1 or y + h > H + 1:
            issues.append("越界 P%d %s" % (n, sh.name))
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP or not sh.has_text_frame:
            continue
        if sh.text_frame.text.strip() and h > 12 and not skip_fit(sh):
            need = need_height(sh, w)
            if need > h - 4:
                issues.append("文字溢出 P%d %s 需%.0fpt 有%.0fpt «%s»" % (
                    n, sh.name, need, h, sh.text_frame.text[:20].replace("\n", "|")))
    for i, a in enumerate(rects):
        for b in rects[i + 1:]:
            if a[5] or b[5] or contains(a, b) or contains(b, a):
                continue
            ox = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            oy = min(a[2] + a[4], b[2] + b[4]) - max(a[2], b[2])
            if ox > 0.6 and oy > 0.6:
                issues.append("重叠 P%d %s × %s (%.0f×%.0fpt)" % (n, a[0], b[0], ox, oy))
    return issues


def main():
    prs = Presentation(sys.argv[1])
    pages = {int(x) for x in sys.argv[2:]}
    W, H = prs.slide_width / PT, prs.slide_height / PT
    issues = []
    for n, sl in enumerate(prs.slides, start=1):
        if not pages or n in pages:
            issues += check_slide(n, sl, W, H)
    for s in issues:
        print(s)
    kinds = [s.split()[0] for s in issues]
    print("\n越界 %d  重叠 %d  文字溢出 %d" % (kinds.count("越界"), kinds.count("重叠"), kinds.count("文字溢出")))
    print("重叠多为有意叠放（图标压横条、文字压底框），溢出可能误报，以渲染图为准。")


if __name__ == "__main__":
    main()
